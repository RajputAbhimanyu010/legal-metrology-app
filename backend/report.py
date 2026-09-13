"""Evidence-first PDF and editable Word report generators."""

import json
import os
from html import escape

from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, KeepTogether
)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


try:
    pdfmetrics.registerFont(TTFont("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    PDF_FONT = "DejaVuSans"
except Exception:
    PDF_FONT = "Helvetica"


def _load_json(value, default):
    if not value:
        return default
    try:
        return json.loads(value) if isinstance(value, str) else value
    except Exception:
        return default


def _recommendations(checks):
    recs = []
    mapping = {
        "Manufacturer Info": "Add a clear manufacturer/packer/importer name and address.",
        "Net Quantity": "Declare the net quantity clearly with the correct unit.",
        "MRP": "Ensure the MRP is clearly visible and includes the required tax declaration.",
        "Manufacture Date": "Print the month/year of manufacture in a clear, readable form.",
        "Consumer Care": "Add complete consumer-care contact details, including phone and email.",
        "Country of Origin": "For imported goods, clearly declare the country of origin.",
        "Font Size": "Increase declaration text size/readability where required.",
    }
    for c in checks:
        if c.get("status") in ("FAIL", "WARN") and c.get("field") in mapping:
            recs.append(mapping[c["field"]])
    return recs or ["No immediate corrective action identified by the deterministic checks."]


def _ai_summary(ai_analysis):
    if not ai_analysis:
        return None, []
    analysis = ai_analysis.get("analysis") if isinstance(ai_analysis, dict) else None
    if not isinstance(analysis, dict):
        return None, []
    summary = analysis.get("product_name") or analysis.get("brand")
    issues = analysis.get("possible_issues") or []
    return summary, issues if isinstance(issues, list) else [str(issues)]


def generate_pdf_report(scan_row: dict, output_dir: str) -> str:
    os.makedirs(output_dir, exist_ok=True)
    file_name = f"report_{scan_row['id']}.pdf"
    file_path = os.path.join(output_dir, file_name)

    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName=PDF_FONT, fontSize=20, leading=24,
                           alignment=TA_CENTER, textColor=colors.HexColor("#123f7a"), spaceAfter=8)
    heading = ParagraphStyle("ReportHeading", parent=styles["Heading2"], fontName=PDF_FONT, fontSize=12, leading=15,
                             textColor=colors.HexColor("#123f7a"), spaceBefore=8, spaceAfter=5)
    small = ParagraphStyle("Small", parent=styles["Normal"], fontName=PDF_FONT, fontSize=8.5, leading=11)

    doc = SimpleDocTemplate(file_path, pagesize=A4, rightMargin=12*mm, leftMargin=12*mm,
                            topMargin=12*mm, bottomMargin=12*mm)
    elements = [Paragraph("Legal Metrology Compliance Checker", title),
                Paragraph("Packaged Commodities Compliance Report", ParagraphStyle("Subtitle", parent=styles["Heading3"], fontName=PDF_FONT)), Spacer(1, 6)]

    score = scan_row.get("compliance_score")
    points_earned = scan_row.get("points_earned")
    max_points = scan_row.get("max_points")
    score_text = "N/A" if score is None else f"{score}%"
    points_text = "N/A" if points_earned is None or max_points is None else f"{points_earned} / {max_points}"

    info = [
        [Paragraph("Report ID", small), Paragraph(str(scan_row["id"]), small), Paragraph("Overall Result", small), Paragraph(str(scan_row.get("overall_status", "N/A")), small)],
        [Paragraph("Product Image", small), Paragraph(escape(scan_row.get("image_name", "")), small), Paragraph("Points Scored", small), Paragraph(points_text, small)],
        [Paragraph("Scan Time", small), Paragraph(escape(str(scan_row.get("scan_time", ""))), small), Paragraph("Compliance Score", small), Paragraph(score_text, small)],
        [Paragraph("Scanned By", small), Paragraph(escape(str(scan_row.get("scanned_by") or "N/A")), small), Paragraph("AI Support", small), Paragraph(escape(str(scan_row.get("ai_provider") or "Not enabled")), small)],
    ]
    info_table = Table(info, colWidths=[27*mm, 68*mm, 32*mm, 55*mm])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#eaf2fb")),
        ("BACKGROUND", (2,0), (2,-1), colors.HexColor("#eaf2fb")),
        ("BOX", (0,0), (-1,-1), 0.5, colors.HexColor("#b9cde4")),
        ("INNERGRID", (0,0), (-1,-1), 0.25, colors.HexColor("#d5e1ef")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    elements += [info_table, Spacer(1, 8)]

    image_path = scan_row.get("image_path")
    if image_path and os.path.exists(image_path):
        elements.append(Paragraph("1. Test / Evidence Image", heading))
        img = Image(image_path)
        max_w, max_h = 176*mm, 92*mm
        scale = min(max_w / img.imageWidth, max_h / img.imageHeight, 1)
        img.drawWidth = img.imageWidth * scale
        img.drawHeight = img.imageHeight * scale
        elements += [img, Paragraph("Original uploaded image used for this compliance report.", small), Spacer(1, 7)]

    checks = _load_json(scan_row.get("checks_json"), [])
    elements.append(Paragraph("2. Compliance Analysis", heading))
    table_data = [[Paragraph("S.No.", small), Paragraph("Requirement", small), Paragraph("Status", small), Paragraph("Points", small), Paragraph("Remarks", small)]]
    for i, c in enumerate(checks, 1):
        status = c.get("status", "INFO")
        points = "1" if status == "PASS" else ("0.5" if status == "WARN" else ("0" if status == "FAIL" else "—"))
        table_data.append([str(i), Paragraph(escape(c.get("field", "")), small), Paragraph(status, small), points, Paragraph(escape(c.get("detail", "")), small)])
    tbl = Table(table_data, colWidths=[10*mm, 42*mm, 22*mm, 18*mm, 90*mm], repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#123f7a")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#b8c7d8")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    elements += [tbl, Spacer(1, 7)]

    elements.append(Paragraph("3. Score & Decision", heading))
    score_box = Table([[Paragraph(f"<b>Overall Result:</b> {escape(str(scan_row.get('overall_status','N/A')))}", small),
                        Paragraph(f"<b>Total Score:</b> {score_text}", small),
                        Paragraph(f"<b>Points:</b> {points_text}", small)]], colWidths=[60*mm, 60*mm, 62*mm])
    score_box.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#eef9f2")),
                                   ("BOX", (0,0), (-1,-1), 0.6, colors.HexColor("#9bd4b1")),
                                   ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8)]))
    elements.append(score_box)

    extracted = scan_row.get("extracted_text") or ""
    if extracted:
        elements.append(Paragraph("4. Extracted Text (OCR + AI)", heading))
        elements.append(Paragraph(escape(extracted).replace("\n", "<br/>"), small))

    ai_analysis = _load_json(scan_row.get("ai_analysis_json"), {})
    ai_summary, ai_issues = _ai_summary(ai_analysis)
    if ai_analysis:
        elements.append(Paragraph("5. AI Visual Review", heading))
        provider = ai_analysis.get("provider_used") or ", ".join(ai_analysis.get("providers", [])) or "AI enabled"
        elements.append(Paragraph(f"Provider: {escape(provider)}", small))
        if ai_summary:
            elements.append(Paragraph(f"AI-observed product/brand: {escape(str(ai_summary))}", small))
        for issue in ai_issues[:8]:
            elements.append(Paragraph(f"• {escape(str(issue))}", small))
        elements.append(Spacer(1, 5))

    elements.append(Paragraph("6. Recommendations", heading))
    for rec in _recommendations(checks):
        elements.append(Paragraph(f"• {escape(rec)}", small))

    elements.append(Spacer(1, 12))
    elements.append(Paragraph("Audit note: AI output is advisory. Final compliance status and points are produced by the deterministic rules engine configured in this application.", small))
    doc.build(elements)
    return file_path


def generate_docx_report(scan_row: dict, output_dir: str) -> str:
    os.makedirs(output_dir, exist_ok=True)
    file_name = f"report_{scan_row['id']}.docx"
    file_path = os.path.join(output_dir, file_name)
    doc = Document()
    title = doc.add_heading("Legal Metrology Compliance Checker", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Packaged Commodities Compliance Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    score = scan_row.get("compliance_score")
    points_earned = scan_row.get("points_earned")
    max_points = scan_row.get("max_points")
    doc.add_paragraph(f"Report ID: {scan_row['id']}")
    doc.add_paragraph(f"Product Image: {scan_row.get('image_name', '')}")
    doc.add_paragraph(f"Scan Time: {scan_row.get('scan_time', '')}")
    doc.add_paragraph(f"Scanned By: {scan_row.get('scanned_by') or 'N/A'}")
    doc.add_paragraph(f"Overall Status: {scan_row.get('overall_status', 'N/A')}")
    doc.add_paragraph(f"Compliance Score: {'N/A' if score is None else str(score) + '%'}")
    doc.add_paragraph(f"Points Scored: {'N/A' if points_earned is None else f'{points_earned} / {max_points}'}")

    image_path = scan_row.get("image_path")
    if image_path and os.path.exists(image_path):
        doc.add_heading("1. Test / Evidence Image", level=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(image_path, width=Inches(5.9))
        cap = doc.add_paragraph("Original uploaded image used for this compliance report.")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    checks = _load_json(scan_row.get("checks_json"), [])
    doc.add_heading("2. Compliance Analysis", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    headers = ["S.No.", "Requirement", "Status", "Points", "Remarks"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    status_colors = {"PASS": RGBColor(0x1e,0x7e,0x34), "FAIL": RGBColor(0xc0,0x39,0x2b), "WARN": RGBColor(0xb9,0x77,0x0e), "INFO": RGBColor(0x5d,0x6d,0x7e)}
    for i, c in enumerate(checks, 1):
        row = table.add_row().cells
        status = c.get("status", "INFO")
        pts = "1" if status == "PASS" else ("0.5" if status == "WARN" else ("0" if status == "FAIL" else "—"))
        values = [str(i), c.get("field", ""), status, pts, c.get("detail", "")]
        for cell, text in zip(row, values):
            cell.text = str(text)
        run = row[2].paragraphs[0].runs[0]
        run.font.bold = True
        if status in status_colors:
            run.font.color.rgb = status_colors[status]

    doc.add_heading("3. Score & Decision", level=2)
    doc.add_paragraph(f"Overall Result: {scan_row.get('overall_status', 'N/A')}")
    doc.add_paragraph(f"Total Score: {'N/A' if score is None else str(score) + '%'}")
    doc.add_paragraph(f"Points: {'N/A' if points_earned is None else f'{points_earned} / {max_points}'}")

    extracted = scan_row.get("extracted_text") or ""
    if extracted:
        doc.add_heading("4. Extracted Text (OCR + AI)", level=2)
        doc.add_paragraph(extracted)

    ai_analysis = _load_json(scan_row.get("ai_analysis_json"), {})
    ai_summary, ai_issues = _ai_summary(ai_analysis)
    if ai_analysis:
        doc.add_heading("5. AI Visual Review", level=2)
        provider = ai_analysis.get("provider_used") or ", ".join(ai_analysis.get("providers", [])) or "AI enabled"
        doc.add_paragraph(f"Provider: {provider}")
        if ai_summary:
            doc.add_paragraph(f"AI-observed product/brand: {ai_summary}")
        for issue in ai_issues[:8]:
            doc.add_paragraph(str(issue), style="List Bullet")

    doc.add_heading("6. Recommendations", level=2)
    for rec in _recommendations(checks):
        doc.add_paragraph(rec, style="List Bullet")

    doc.add_paragraph("Audit note: AI output is advisory. Final compliance status and points are produced by the deterministic rules engine configured in this application.")
    doc.add_heading("Officer Notes (editable)", level=2)
    doc.add_paragraph("\n\n____________________________________________________________")
    doc.save(file_path)
    return file_path
